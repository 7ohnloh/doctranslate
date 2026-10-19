# Original version (v1), recovered from the shipped PyInstaller build.
# Kept for before/after comparison only. The redesigned app lives in /translator.
import os
import threading
import tkinter as tk
from tkinter import filedialog, ttk, messagebox
from pptx import Presentation
from docx import Document
from googletrans import Translator, LANGUAGES
import pytesseract
from PIL import Image
import io


def get_downloads_folder():
    home = os.path.expanduser("~")
    downloads = os.path.join(home, "Downloads")
    return downloads


class TranslatorEngine:
    def __init__(self, source_lang="en", target_lang="zh-cn", enable_ocr=False, progress_callback=None):
        self.source_lang = source_lang
        self.target_lang = target_lang
        self.enable_ocr = enable_ocr
        self.progress_callback = progress_callback
        self.translator = Translator()
        self.translator.raise_Exception = False
        self.translation_cache = {}

    def update_progress(self, value):
        if self.progress_callback:
            self.progress_callback(value)

    def translate_text(self, text, retries=3):
        text = text.strip()
        if not text:
            return text
        if text in self.translation_cache:
            return self.translation_cache[text]
        for attempt in range(retries):
            try:
                try:
                    detected = self.translator.detect(text).lang
                except:
                    detected = self.source_lang
                if detected == self.target_lang:
                    self.translation_cache[text] = text
                    return text
                translated = self.translator.translate(text, src=self.source_lang, dest=self.target_lang).text
                self.translation_cache[text] = translated
                return translated
            except Exception:
                import time
                time.sleep(0.4)
        return text

    def translate_docx(self, filepath):
        doc = Document(filepath)
        blocks = []
        for p in doc.paragraphs:
            blocks.append(p)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        blocks.append(p)
        total = len(blocks)
        for i, paragraph in enumerate(blocks, start=1):
            if paragraph.text.strip():
                paragraph.text = self.translate_text(paragraph.text)
            self.update_progress(i / total * 100)
        out_name = f"translated_{self.target_lang}.docx"
        out_path = os.path.join(get_downloads_folder(), out_name)
        doc.save(out_path)
        return out_path

    def translate_pptx(self, filepath):
        prs = Presentation(filepath)
        text_items = []

        def collect(slide):
            for shape in slide.shapes:
                if hasattr(shape, "has_text_frame") and shape.has_text_frame and shape.text.strip():
                    text_items.append(shape)
                elif shape.shape_type == 6:  # group
                    for s in shape.shapes:
                        if hasattr(s, "has_text_frame") and s.has_text_frame and s.text.strip():
                            text_items.append(s)
                elif hasattr(shape, "has_table") and shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                text_items.append(cell)

        for slide in prs.slides:
            collect(slide)
        total = len(text_items)
        for i, shape in enumerate(text_items, start=1):
            text = shape.text.strip()
            if text:
                shape.text = self.translate_text(text)
            self.update_progress(i / total * 100)

        if self.enable_ocr:
            for slide in prs.slides:
                for shape in slide.shapes:
                    try:
                        if shape.shape_type == 13:  # picture
                            image = shape.image
                            img = Image.open(io.BytesIO(image.blob))
                            text = pytesseract.image_to_string(img, lang="eng").strip()
                            if text:
                                self.translate_text(text)  # result was discarded
                    except:
                        pass

        out_name = f"translated_{self.target_lang}.pptx"
        out_path = os.path.join(get_downloads_folder(), out_name)
        prs.save(out_path)
        return out_path


class TranslationApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Internal File Translator Tool")
        self.root.geometry("500x420")
        self.root.resizable(False, False)
        self.filepath = None

        tk.Label(root, text="Select File Type:", font=("Segoe UI", 11)).pack(pady=5)
        self.file_type_var = tk.StringVar()
        self.file_type_dropdown = ttk.Combobox(root, textvariable=self.file_type_var, values=["DOCX", "PPTX"], state="readonly")
        self.file_type_dropdown.pack(pady=5)

        tk.Button(root, text="Choose File", command=self.choose_file).pack(pady=8)
        self.file_label = tk.Label(root, text="No file selected", fg="gray")
        self.file_label.pack()

        tk.Label(root, text="Select Target Language:", font=("Segoe UI", 11)).pack(pady=10)
        self.lang_var = tk.StringVar()
        lang_list = [f"{name.title()} ({code})" for code, name in LANGUAGES.items()]
        self.lang_dropdown = ttk.Combobox(root, textvariable=self.lang_var, values=lang_list)
        self.lang_dropdown.pack(pady=5)

        self.ocr_var = tk.BooleanVar()
        tk.Checkbutton(root, text="Enable Image Detection (OCR)", variable=self.ocr_var).pack(pady=10)

        tk.Label(root, text="Progress:", font=("Segoe UI", 11)).pack()
        self.progress = ttk.Progressbar(root, length=350, mode="determinate")
        self.progress.pack(pady=10)

        tk.Button(root, text="Start Translation", font=("Segoe UI", 11, "bold"), command=self.start_translation_thread).pack(pady=10)

    def choose_file(self):
        file = filedialog.askopenfilename(filetypes=[("Documents", "*.docx *.pptx")])
        if file:
            self.filepath = file
            self.file_label.config(text=os.path.basename(file))

    def start_translation_thread(self):
        if not self.filepath:
            messagebox.showerror("Error", "Please choose a file.")
            return
        if not self.lang_var.get():
            messagebox.showerror("Error", "Please choose a target language.")
            return
        lang = self.extract_lang_code(self.lang_var.get())
        if not lang:
            messagebox.showerror("Error", "Invalid language selection.")
            return
        thread = threading.Thread(target=self.translate_file, args=(lang,))
        thread.start()

    def translate_file(self, target_lang):
        try:
            self.progress["value"] = 0
            filetype = self.file_type_var.get()
            engine = TranslatorEngine(target_lang=target_lang, enable_ocr=self.ocr_var.get(), progress_callback=self.progress_callback)
            if filetype == "DOCX":
                out_file = engine.translate_docx(self.filepath)
            else:
                out_file = engine.translate_pptx(self.filepath)
            messagebox.showinfo("Success!", f"File saved to:\n{out_file}")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def progress_callback(self, value):
        self.progress["value"] = value

    def extract_lang_code(self, text):
        if "(" in text and ")" in text:
            return text.split("(")[-1].replace(")", "").strip()
        return None


if __name__ == "__main__":
    root = tk.Tk()
    app = TranslationApp(root)
    root.mainloop()
