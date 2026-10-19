"""Translation engine for .docx and .pptx files.

UI-independent: the app talks to it only through `on_progress` and `cancel_event`,
so it can also be used from a script or tested on its own.
"""
from __future__ import annotations

import io
import os
import shutil
import threading
import time
from dataclasses import dataclass
from typing import Callable, Iterable, List, Optional

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SUPPORTED_EXTENSIONS = (".docx", ".pptx")


class TranslationCancelled(Exception):
    pass


@dataclass
class Progress:
    done: int
    total: int
    message: str

    @property
    def fraction(self) -> float:
        return self.done / self.total if self.total else 1.0


@dataclass
class TranslationResult:
    output_path: str
    segments: int
    images_read: int = 0
    failed: int = 0  # blocks left in the original language after retries


def file_kind(path: str) -> Optional[str]:
    ext = os.path.splitext(path)[1].lower()
    return ext[1:].upper() if ext in SUPPORTED_EXTENSIONS else None


# --------------------------------------------------------------------------- OCR

_WINDOWS_TESSERACT = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def find_tesseract() -> Optional[str]:
    """Return the Tesseract binary path, or None if OCR isn't available."""
    path = shutil.which("tesseract")
    if path:
        return path
    if os.name == "nt" and os.path.exists(_WINDOWS_TESSERACT):
        return _WINDOWS_TESSERACT
    return None


# ------------------------------------------------------------ text replacement

def _replace_paragraph_text(paragraph, new_text: str) -> None:
    """Swap a paragraph's text but keep the formatting of its first run.

    Setting `paragraph.text` directly would drop every run and reset font, size,
    bold and colour, which is what v1 did.
    """
    runs = [r for r in paragraph.runs if r.text]
    if not runs:
        paragraph.text = new_text
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def _unique(paragraphs: Iterable) -> List:
    # Merged table cells return the same paragraph more than once.
    seen, out = set(), []
    for p in paragraphs:
        key = id(p._p)
        if key not in seen and p.text.strip():
            seen.add(key)
            out.append(p)
    return out


def _docx_paragraphs(doc) -> List:
    def from_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
                    yield from from_tables(cell.tables)

    def from_container(container):
        yield from container.paragraphs
        yield from from_tables(container.tables)

    items = list(from_container(doc))
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                items.extend(from_container(part))
    return _unique(items)


def _pptx_paragraphs(shapes) -> List:
    items = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            items.extend(_pptx_paragraphs(shape.shapes))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            items.extend(shape.text_frame.paragraphs)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    items.extend(cell.text_frame.paragraphs)
    return items


def _pptx_pictures(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape


def count_segments(path: str) -> int:
    """Number of text blocks that will be translated. Used for the file preview."""
    kind = file_kind(path)
    if kind == "DOCX":
        return len(_docx_paragraphs(Document(path)))
    if kind == "PPTX":
        prs = Presentation(path)
        return len(_unique(p for s in prs.slides for p in _pptx_paragraphs(s.shapes)))
    raise ValueError("Unsupported file type")


def unique_output_path(folder: str, source_path: str, lang_code: str) -> str:
    stem, ext = os.path.splitext(os.path.basename(source_path))
    candidate = os.path.join(folder, f"{stem} ({lang_code}){ext}")
    n = 2
    while os.path.exists(candidate):
        candidate = os.path.join(folder, f"{stem} ({lang_code}) {n}{ext}")
        n += 1
    return candidate


# ---------------------------------------------------------------------- engine

class TranslatorEngine:
    def __init__(
        self,
        target_lang: str,
        source_lang: str = "auto",
        enable_ocr: bool = False,
        on_progress: Optional[Callable[[Progress], None]] = None,
        cancel_event: Optional[threading.Event] = None,
        translator=None,
    ):
        self.source_lang = source_lang
        self.target_lang = target_lang
        self.enable_ocr = enable_ocr
        self.on_progress = on_progress
        self.cancel_event = cancel_event or threading.Event()
        if translator is None:
            from googletrans import Translator
            translator = Translator()
            translator.raise_Exception = False
        self.translator = translator
        self.cache: dict = {}

    # -- helpers
    def _report(self, done: int, total: int, message: str) -> None:
        if self.on_progress:
            self.on_progress(Progress(done, total, message))

    def _check_cancel(self) -> None:
        if self.cancel_event.is_set():
            raise TranslationCancelled()

    def translate_text(self, text: str, retries: int = 3) -> str:
        stripped = text.strip()
        if not stripped:
            return text
        if stripped in self.cache:
            return self.cache[stripped]
        last_error = None
        for attempt in range(retries):
            try:
                # One request instead of v1's detect() + translate() pair.
                result = self.translator.translate(stripped, src=self.source_lang, dest=self.target_lang)
                self.cache[stripped] = result.text
                return result.text
            except Exception as exc:  # network hiccups, rate limits
                last_error = exc
                time.sleep(0.4 * (attempt + 1))
        raise ConnectionError(
            "Couldn't reach the translation service. Check your internet connection and try again."
        ) from last_error

    def _translate_paragraphs(self, paragraphs: List, label: str) -> int:
        """Translate in place. Returns how many blocks failed and were left as-is."""
        total = len(paragraphs)
        failed = consecutive = 0
        self._report(0, total, f"Preparing {total} {label}…")
        for i, paragraph in enumerate(paragraphs, start=1):
            self._check_cancel()
            try:
                _replace_paragraph_text(paragraph, self.translate_text(paragraph.text))
                consecutive = 0
            except ConnectionError:
                failed += 1
                consecutive += 1
                if consecutive >= 3:  # the service is down, not a one-off
                    raise
            self._report(i, total, f"Translating {i} of {total} {label}")
        return failed

    # -- public API
    def translate_file(self, path: str, output_folder: str) -> TranslationResult:
        kind = file_kind(path)
        if kind == "DOCX":
            return self.translate_docx(path, output_folder)
        if kind == "PPTX":
            return self.translate_pptx(path, output_folder)
        raise ValueError("Only .docx and .pptx files are supported.")

    def translate_docx(self, path: str, output_folder: str) -> TranslationResult:
        doc = Document(path)
        paragraphs = _docx_paragraphs(doc)
        failed = self._translate_paragraphs(paragraphs, "paragraphs")
        out_path = unique_output_path(output_folder, path, self.target_lang)
        doc.save(out_path)
        return TranslationResult(out_path, len(paragraphs), failed=failed)

    def translate_pptx(self, path: str, output_folder: str) -> TranslationResult:
        prs = Presentation(path)
        paragraphs = _unique(p for slide in prs.slides for p in _pptx_paragraphs(slide.shapes))
        failed = self._translate_paragraphs(paragraphs, "text blocks")

        images_read = 0
        if self.enable_ocr:
            images_read = self._ocr_slides(prs)

        out_path = unique_output_path(output_folder, path, self.target_lang)
        prs.save(out_path)
        return TranslationResult(out_path, len(paragraphs), images_read, failed)

    def _ocr_slides(self, prs) -> int:
        """Read text inside pictures and add the translation to the slide notes.

        v1 ran OCR but threw the result away. Rewriting text inside an image isn't
        possible without editing pixels, so speaker notes are the place for it.
        """
        import pytesseract
        from PIL import Image

        binary = find_tesseract()
        if not binary:
            return 0
        pytesseract.pytesseract.tesseract_cmd = binary

        slides = list(prs.slides)
        read = 0
        for n, slide in enumerate(slides, start=1):
            self._check_cancel()
            self._report(n, len(slides), f"Reading images on slide {n} of {len(slides)}")
            found = []
            for picture in _pptx_pictures(slide.shapes):
                try:
                    text = pytesseract.image_to_string(Image.open(io.BytesIO(picture.image.blob))).strip()
                except Exception:
                    continue
                if text:
                    try:
                        found.append(self.translate_text(text))
                        read += 1
                    except ConnectionError:
                        pass
            if found:
                notes = slide.notes_slide.notes_text_frame
                block = "Translated image text:\n" + "\n\n".join(found)
                notes.text = f"{notes.text}\n\n{block}" if notes.text.strip() else block
        return read
